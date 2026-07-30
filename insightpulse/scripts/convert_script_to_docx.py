import docx
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def main():
    doc = Document()
    
    # Page setup
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Document title
    title = doc.add_heading('Video Demo Script: InsightPulse', level=1)
    
    # Introductory text
    doc.add_paragraph(
        "This script is structured to help you record a conversational, authentic, and highly technical "
        "10–15 minute video presentation of your internship project. It is structured around the 11 slides "
        "in your presentation and guides you through the live demonstration and code walkthrough."
    )
    
    # Technical Setup Heading
    doc.add_heading('Technical Setup Before Recording', level=2)
    
    setup_items = [
        "Screen Resolution: Set your monitor scale so text is easily readable.",
        "Code Editor: Increase font size in VS Code (zoom in twice).",
        "Tabs Ready in Browser:\n  - Tab 1: Slide 1 of the PowerPoint deck (InsightPulse_Presentation.pptx) open in slideshow mode.\n  - Tab 2: The InsightPulse live app dashboard (http://localhost:5001 or https://insightpulse.skill-pulse.io).\n  - Tab 3: The database console or Zoho CRM sandbox (if demonstrating live sync).",
        "Code Files Open in VS Code:\n  - server/services/surveyResponseAnalysisService.ts (LLM Analysis & Prompt anatomy)\n  - server/routes/voiceRoutes.ts (VAPI webhooks and callbacks)\n  - server/routes/whatsappRoutes.ts (WhatsApp state machine)"
    ]
    
    for i, item in enumerate(setup_items, 1):
        doc.add_paragraph(f"{i}. {item}")
        
    # Timeline Heading
    doc.add_heading('Walkthrough Timeline & Script', level=2)
    
    # Define table data
    table_data = [
        ("Section / Time", "Visual Target (On Screen)", "Spoken Narration (Conversational Guidance)"),
        (
            "Section 1: Intro\n0:00 - 1:00",
            "Slide 1: Title Slide\n(Show the official cover slide with logos and student details).",
            "Hello everyone, my name is Bendi Karthikeya, and I am a student in the BSc. Honours in Data Science and Artificial Intelligence programme at IIT Guwahati. Today, I am presenting my Internship-II project submission for course DA 378.\n\nThe project is called InsightPulse: An AI-Powered, Multi-Channel Customer Feedback and Unified Analytics Platform. I developed this project during my remote internship as a Full Stack Generative AI Intern at FIT - FIRST Information Technology in Muscat, Oman."
        ),
        (
            "Section 2: Context\n1:00 - 2:00",
            "Slide 2: Introduction & Internship\n(Show the overview bullet points).",
            "To give you some context on where I interned: FIT is a software solutions provider building enterprise-grade applications for regional clients. My work on InsightPulse focused on combining React, Express, PostgreSQL, and LLMs to create a unified customer feedback pipeline.\n\nWhile my broader internship allowed me to work with a range of tools like CRM OAuth, voice APIs, and WhatsApp webhook integration, InsightPulse brings all of these together into a single, cohesive, production-ready system."
        ),
        (
            "Section 3: The Problem\n2:00 - 3:15",
            "Slide 3: The Problem & Why It Matters\n(Show the slide highlighting data fragmentation).",
            "Let's discuss the problem. Enterprises collect feedback from various places—web forms, emails, WhatsApp messages, and voice calls. Traditionally, these exist in separate data silos. Voice transcripts are rarely analyzed at scale because manual review is too slow and expensive. When businesses do collect written text, categorizing comments and determining specific customer emotions is slow and subjective.\n\nOur goal was to build a system where all channels write to a single schema, and we use a resilient LLM pipeline to analyze sentiment and emotions in real-time."
        ),
        (
            "Section 4: Stack\n3:15 - 4:15",
            "Slide 4: Tools & Technologies\n(Show the Tech Stack table).",
            "Here is the technology stack I used to build the platform. The frontend is built on React 18 with TypeScript and Tailwind CSS, using TanStack Query for real-time polling. The backend runs on Node.js and Express, with Neon Serverless PostgreSQL as the database and Drizzle ORM for query management.\n\nFor the AI and channel integrations, we use VAPI for voice assistant automation, the Meta WhatsApp Cloud API for messaging, and OpenRouter to query LLMs—specifically using DeepSeek-Chat as our primary model and Gemini-2.0-Flash as our automatic failover."
        ),
        (
            "Section 5: Ingested Inputs\n4:15 - 5:15",
            "Slide 5: System Ingested Input\n(Show the slide detailing the feedback database).",
            "Because this is an engineering internship project focused on system implementation, we do not work with static training/testing datasets. Instead, InsightPulse handles a dynamic stream of customer feedback. The data comprises scores for CSAT, NPS, Customer Effort Score (CES), and the Emotional Value Index (EVI), coupled with open-ended user comments.\n\nOur preprocessing normalizes these varied inputs: VAPI provides structured call metrics via webhooks, and WhatsApp/web messages are mapped to our database using a clean relational model before being sent to the AI sentiment pipeline."
        ),
        (
            "Section 6: Ingestion\n5:15 - 6:30",
            "Slide 6: How It Works\n(Show the visual pipeline diagram).",
            "This diagram shows the system architecture. When a survey response is submitted, it enters the Express backend. If it's a web or WhatsApp submission, the raw text is sent asynchronously to the OpenRouter pipeline. The LLM extracts sentiment, primary categories, and 8 basic Plutchik emotions.\n\nIf it's a voice call, VAPI conducts the conversation and sends the structured scores directly. Everything is saved to PostgreSQL, which triggers a background recalculation of department NPS. The frontend dashboard then polls these tables every 30 seconds."
        ),
        (
            "Section 7: Live Demo\n6:30 - 9:30",
            "Web Browser: Live App\n(Navigate to your running dashboard, switch tabs, show Recharts visualizations, CSAT charts, Zoho contact sync, and trigger a test survey).",
            "Let's look at the live platform. Here is the Overview Tab showing EVI, NPS, CSAT, and CES with period-over-period trends. If we go to the CSAT/NPS Tab, we see a breakdown of promoters and satisfaction distributions. The Customer Journey Tab tracks metrics across touchpoints, and the Trends Tab plots Plutchik emotions over time.\n\nUnder Insights, we see automated action plans generated by the LLM. In the CRM settings, I can connect to Zoho CRM via OAuth, fetch contacts, and select contacts to receive surveys via WhatsApp or Voice. Once submitted, the scores write back to Zoho automatically."
        ),
        (
            "Section 8: Code Walkthrough\n9:30 - 12:30",
            "VS Code: Editor\n(Show the code files, highlighting prompt engineering, the failover chain, and webhook handlers).",
            "Now let's look at the code. Opening surveyResponseAnalysisService.ts, you can see the prompt engineering in buildAnalysisPrompt(). It sets the system role, passes the survey text, and defines the Plutchik emotion scales and categories. We enforce structured output using JSON mode with a low temperature of 0.3.\n\nNotice the failover logic: we call DeepSeek-Chat first, and if it fails or returns malformed text, our catch block transparently retries using Gemini-2.0-Flash.\n\nIn whatsappRoutes.ts, we maintain the session state in the database using the wati_sessions table, sending consent requests and delivering questions sequentially. Over in voiceRoutes.ts, you can see the webhook endpoint that receives VAPI call completed payloads and maps them to our unified schema."
        ),
        (
            "Section 9: Academic Connection\n12:30 - 13:30",
            "Slide 8: Academic Foundation\n(Show Table III mapping LLM course concepts).",
            "This AI pipeline directly applies concepts from the Large Language Model Application Engineering (DA 3022) course. The prompt design follows the six-part prompt anatomy. We use chain-of-thought instructions by telling the model to extract facts before scoring. We enforce structured outputs using JSON mode and Zod schemas, guarantee reliability with a low temperature, and protect against outages by setting up a resilient model failover chain."
        ),
        (
            "Section 10: Conclusion\n13:30 - 14:30",
            "Slide 11: Wrap Up\n(Show the final slide with summary, links, and future work).",
            "To wrap up: InsightPulse successfully consolidates Web, WhatsApp, and Voice surveys into a single dashboard. For future work, we plan to use LangGraph for conversational multi-turn voice surveys, implement semantic clustering to group customer complaints, and create automated email notifications for negative customer sentiment.\n\nThe source code is available on GitHub at the link shown. Thank you for listening, and I am happy to answer any questions!"
        )
    ]
    
    # Create Table
    table = doc.add_table(rows=len(table_data), cols=3)
    table.style = 'Table Grid'
    
    # Width variables
    widths = [Inches(1.5), Inches(1.8), Inches(3.2)]
    
    # Set headers
    hdr_cells = table.rows[0].cells
    for j in range(3):
        hdr_cells[j].text = table_data[0][j]
        hdr_cells[j].width = widths[j]
        # Make bold
        for p in hdr_cells[j].paragraphs:
            p.runs[0].font.bold = True
            
    # Set data rows
    for i, row in enumerate(table_data[1:], 1):
        row_cells = table.rows[i].cells
        for j in range(3):
            row_cells[j].text = row[j]
            row_cells[j].width = widths[j]
            
    # Pro-tips Heading
    doc.add_heading('Pro-Tips for Recording', level=2)
    
    tips = [
        "Microphone: Use a good quality headset mic. Record a 10-second test first and play it back to check audio levels.",
        "Pacing: Don't rush through the live demo. Move your mouse cursor slowly to guide the viewer's eyes.",
        "Mistakes: If you make a mistake, don't worry! Just pause for 2 seconds, correct yourself, and continue. You don't need to stop and start the whole recording over.",
        "YouTube Upload: Set the video visibility to Unlisted or Public so your instructors can view it. Use Slide 1 of the PowerPoint deck as the thumbnail."
    ]
    
    for tip in tips:
        doc.add_paragraph(f"• {tip}")
        
    doc.save('docs/iitg/InsightPulse_Video_Script.docx')
    print("Docx file generated successfully!")

if __name__ == "__main__":
    main()
