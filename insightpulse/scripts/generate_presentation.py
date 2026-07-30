import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def replace_text_recursively(shape, replacements):
    """Recursively replaces text in groups or nested shapes."""
    if shape.shape_type == 6: # Group shape
        for subshape in shape.shapes:
            replace_text_recursively(subshape, replacements)
    elif shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            orig_text = paragraph.text
            new_text = orig_text
            replaced = False
            for old, new in replacements.items():
                if old in new_text:
                    new_text = new_text.replace(old, new)
                    replaced = True
            if replaced:
                paragraph.text = new_text

def apply_card_styling(shape, bg_color=RGBColor(250, 252, 254), border_color=RGBColor(226, 232, 240)):
    """Styles a textbox/shape as a beautiful container card with background and border."""
    try:
        # Enable solid fill
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        
        # Enable border line
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        # Add internal margins (padding) so text does not touch the borders
        shape.text_frame.margin_left = Inches(0.4)
        shape.text_frame.margin_right = Inches(0.4)
        shape.text_frame.margin_top = Inches(0.4)
        shape.text_frame.margin_bottom = Inches(0.4)
    except Exception as e:
        pass

def format_paragraphs_in_shape(shape, title_font="Trebuchet MS", body_font="Arial"):
    """Applies modern fonts and hierarchy styles to paragraphs inside a text frame."""
    if not shape.has_text_frame:
        return
        
    for p in shape.text_frame.paragraphs:
        stripped = p.text.strip()
        if not stripped:
            continue
            
        # Is it a bullet title?
        is_bullet_title = stripped.startswith("•") or (hasattr(p, "font") and p.font.bold)
        
        # Enforce properties run-by-run
        for run in p.runs:
            run.font.name = title_font if is_bullet_title else body_font
            run.font.size = Pt(14.5) if is_bullet_title else Pt(13)
            run.font.bold = is_bullet_title
            
            if is_bullet_title:
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate 900
            else:
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69) # Slate 600
                
        # Spacing and levels
        if is_bullet_title:
            p.space_before = Pt(8)
            p.space_after = Pt(2)
        else:
            p.space_before = Pt(0)
            p.space_after = Pt(4)
            p.level = 1 # Indent bullet descriptions

def main():
    src_path = "/Users/karthikeya/Downloads/Omen/insightpulse/docs/iitg/PPT.pptx"
    dest_path = "/Users/karthikeya/Downloads/Omen/insightpulse/docs/iitg/InsightPulse_Presentation.pptx"
    
    # Copy file directly to avoid corrupting slide properties or relationships
    if os.path.exists(dest_path):
        os.remove(dest_path)
    shutil.copy(src_path, dest_path)
    
    prs = Presentation(dest_path)
    
    # --- Slide 1 Replacements ---
    slide1_replacements = {
        "DA 377": "DA 378",
        "Internship-I": "Internship-II",
        "VII": "IX", # Trimester 9
        "HCMS Recruitment Suite": "InsightPulse: AI-Powered Multi-Channel Survey & Unified Analytics Platform"
    }
    
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        replace_text_recursively(shape, slide1_replacements)
        
    # Format Slide 1 Title specifically
    for shape in slide1.shapes:
        if shape.has_text_frame and "InsightPulse" in shape.text_frame.text:
            shape.height = Inches(1.6) # Increase height to allow wrapping without cutting
            shape.width = Inches(9.2)  # Increase width slightly to give more room
            for p in shape.text_frame.paragraphs:
                p.font.name = "Trebuchet MS"
                p.font.size = Pt(28) # Clean fit
                p.font.bold = True
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    # --- Slides 2-11 Replacements ---
    footer_repl = {"HCMS Recruitment Suite": "InsightPulse"}
    
    # Slide 2: Introduction
    slide2_repl = {
        "Bendi Karthikeya, BSc (Hons) Data Science & AI, IIT Guwahati (Online). This is my Internship-I (DA 377) submission.":
        "Bendi Karthikeya, BSc (Hons) Data Science & AI, IIT Guwahati (Online). This is my Internship-II (DA 378) submission.",
        "The HCMS Recruitment Suite (Human Capital Management System) is the client project I built at FIT, a full recruitment platform in Django / Python, with an AI CV-screening engine at its core.": 
        "InsightPulse is the client project I built at FIT: an AI-powered multi-channel survey platform collecting feedback via Voice (VAPI), WhatsApp, Web, and Email with real-time unified analytics."
    }
    
    # Slide 3: The Problem
    slide3_repl = {
        "Hiring is process-heavy.": "Feedback collection is fragmented.",
        "One hire crosses 6+ roles (requester, HOD, HR, finance, legal, management) and a dozen documents: requisitions, JDs, scorecards, proposals, offers.":
        "Surveys are distributed across separate, disconnected systems (email, SMS, web, voice) with no unified data structure or central storage.",
        "It runs on e-mail, spreadsheets, wet ink.": "Voice & WhatsApp data is lost.",
        "Screening CVs by hand is slow and inconsistent, and signatures cannot be verified later.":
        "Voice conversations are not transcribed or analyzed, and messaging is managed manually, leading to delayed feedback loops.",
        "Fragmented tools lose the audit trail.": "Analysis lacks depth and speed.",
        "You cannot reliably say who approved what, when, and with which signature, exactly what compliance needs.":
        "Manual categorization of user comments is slow, subjective, and unable to extract emotional context or action items in real-time.",
        "One platform where the whole funnel lives together, CV screening is done by AI, and every decision is recorded for audit.":
        "A unified platform where web, WhatsApp, and AI voice surveys dump into one database, with real-time analytics and LLM-powered emotion scoring."
    }
    
    # Slide 4: Tools & Technologies
    slide4_repl = {
        "This project, HCMS Recruitment Suite": "This project, InsightPulse",
        "Python · Django (backend + web app)": "React 18 · TypeScript (Frontend layout)",
        "PostgreSQL (database)": "Node.js · Express (Backend API server)",
        "LLM via API, Llama-3.1-8B (Groq), DeepSeek (OpenRouter)": "Neon Serverless PostgreSQL (Database storage)",
        "PyPDF2 · python-docx (CV text extraction)": "Drizzle ORM (Database query schema)",
        "DocuSign · Adobe Acrobat Sign (e-signatures)": "VAPI Voice AI & Meta WhatsApp Cloud API",
        "No LangChain / LangGraph, direct API calls": "OpenRouter (DeepSeek/Gemini fallback APIs)",
        "Across the internship, FIT (context)": "Core Features Highlight",
        "React · TypeScript · Node.js · Express": "Real-time TanStack Query dashboard (30s poll)",
        "VAPI Voice AI · WhatsApp Business API": "Plutchik's 8-Emotion LLM scoring model",
        "Retrieval-Augmented Generation (RAG)": "Zoho CRM OAuth contacts integration",
        "OAuth-based CRM integrations": "Dynamic question builder & localized translation",
        "Drizzle ORM · analytics dashboards": "Automated department NPS scoring & action plans"
    }
    
    # Slide 5: The Data
    slide5_repl = {
        "The Data": "System Ingested Input",
        "Not a fixed dataset.": "Real-time feedback stream.",
        "The data is the candidate's own documents, a CV or résumé as a PDF or Word file.":
        "The system does not process static datasets, but handles real-time feedback submissions from active channels.",
        "Paired with the job requirement.": "Quantitative & Qualitative inputs.",
        "Each CV is screened against the text of the role the candidate applied for.":
        "Captures both numerical scores (CSAT 1-5, NPS 0-10, EVI 0-100) and open-ended customer comments.",
        "Preprocessing = text extraction.": "Asynchronous normalization.",
        "I pull raw text from PDFs with PyPDF2 and from Word files with python-docx; that text plus the job requirement is the input to the LLM.":
        "Numerical feedback and raw transcripts are structured dynamically. Raw texts are queued for OpenRouter emotion extraction, then saved in SQL."
    }
    
    # Slide 6: How It Works
    slide6_repl = {
        "How It Works, The Hiring Funnel": "How It Works: Data Pipeline",
        "Django modular monolith; every transition records an audit row":
        "Multi-channel ingestion with unified analytical storage",
        "Status changes only through engine functions, never ad-hoc edits, so the whole funnel is replayable for audit.":
        "Feedback flows from channels -> webhook / API -> OpenRouter analysis -> DB save -> TanStack Query UI updates."
    }
    
    # Slide 7: AI & Voice AI Integration
    slide7_repl = {
        "AI CV Screening Engine": "AI & Voice AI Integration",
        "recruitment/cv_screening_ai.py, a three-tier fallback, so it never hard-fails":
        "Voice calls via VAPI, text analysis via OpenRouter (DeepSeek/Gemini failover)",
        "Engineered prompt, _build_prompt().": "Structured VAPI Callbacks.",
        "Role ('senior HR evaluator') + task + CV/JD context + an 8-criterion rubric (80 pts) + strict JSON schema. The STEP 1 extract → 2 score → 3 total → 4 threshold protocol is chain-of-thought.":
        "VAPI Voice assistant calls customers, asks localized (EN/AR) questions, captures voice responses, extracts structured survey scores, and calls the webhook.",
        "Structured + reproducible.": "OpenRouter LLM Pipeline.",
        "response_format = JSON object forces valid JSON; temperature 0.3 makes the same CV score the same way.":
        "Web and WhatsApp survey comments are sent to OpenRouter. Strict system prompts enforce JSON mode and a temperature of 0.3 ensures consistent evaluation.",
        "Explainable thresholds.": "Sentiment & Plutchik Emotions.",
        "≥70% auto-shortlist, ≥60% interview, <60% reject, mirroring the company's paper evaluation form.":
        "LLM rates 8 Plutchik emotions (0-10), extracts sentiment, EVI score, category/theme, and generates real-time recommendations and action items.",
        "Non-blocking.": "Resilient Failover Chain.",
        "Inference runs in a background thread; the UI shows 'screening…' and polls for the score.":
        "Primary model is DeepSeek-Chat. If the API fails or returns invalid JSON, it transparently falls back to Gemini-2.0-Flash on OpenRouter."
    }
    
    # Slide 8: Academic Foundation
    slide8_repl = {
        "role + task + context + rubric + constraints + schema → _build_prompt()":
        "role + context + 8-emotion rubric + instructions + JSON schema → buildAnalysisPrompt()",
        "the STEP 1 → 2 → 3 → 4 protocol inside the prompt":
        "the STEP 1 parse -> 2 categorise -> 3 evaluate emotions -> 4 recommend actions protocol",
        "response_format = json_object; strict JSON schema":
        "response_format = {\"type\": \"json_object\"} + Zod validation schema",
        "temperature 0.3 for reproducible scores":
        "temperature 0.3 for reproducible sentiment scores",
        "'Return ONLY valid JSON', 'do not speculate'":
        "'Return ONLY valid JSON', 'do not add markdown formatting or extra text'"
    }
    
    # Slide 9: CRM Integration
    slide9_repl = {
        "Sequential Approval Engine": "CRM Integration & Custom Flows",
        "The other core idea, and real e-signatures":
        "Zoho CRM OAuth, WhatsApp Sequences & Dept NPS",
        "Chains are ordered step rows.": "Zoho CRM Integration.",
        "One row per decision point (sequence, approver, status, signature).":
        "OAuth 2.0 flow allows administrators to fetch client contacts from Zoho, trigger multi-channel surveys, and sync responses back to CRM records.",
        "Actionability is derived, not stored.": "WhatsApp Session State.",
        "First pending step is ACTIVE; all later steps WAIT. Out-of-order or duplicate signing is impossible by design.":
        "WatiSessions tracking allows sequential question delivery (checking response data types per step) to avoid multi-message clutter on WhatsApp.",
        "Same engine everywhere.": "Department-wise NPS.",
        "Requisitions, employment proposals, and offer letters all reuse it.":
        "Aggregates Net Promoter Score dynamically for departments, mapping employee feedback back to their respective business units.",
        "Real e-signatures.": "Scheduled Survey Cron.",
        "The offer is signed through DocuSign or Adobe Acrobat Sign, and the signed PDF is stored back on the record.":
        "Custom survey scheduler executes in the background to automatically dispatch pending survey campaigns via email/WhatsApp."
    }
    
    # Slide 10: Outputs & Results
    slide10_repl = {
        "Shown live in the video, walking through the running platform":
        "Demonstrated in the video walk-through of the live platform",
        "Upload a CV.": "Trigger Outbound Voice AI.",
        "The candidate row shows 'screening…', then flips to a score with the 8-criterion breakdown and matched vs. missing skills.":
        "Initiating a VAPI call dials the target, conducts the survey, and displays the transcript and structured output in the dashboard.",
        "Ranked pipeline.": "Unified Real-Time Dashboard.",
        "Candidates are ordered by match score, so the strongest are on top.":
        "React UI displays EVI, NPS, CSAT, and CES scores with period-over-period trends and 30s auto-refresh.",
        "Interviews → proposal.": "LLM Emotion Radar.",
        "Rounds are scored; when complete, the candidate becomes an employment proposal.":
        "Sentiment classification and Plutchik's 8-emotion distribution charts dynamically render customer emotional profiles.",
        "E-signed offer.": "Zoho CRM Synchronization.",
        "The approval chain shows one active signer at a time, ending in an offer letter signed through DocuSign / Adobe.":
        "Surveys triggered for CRM contacts correctly update fields, enabling automated customer relationship management loops."
    }
    
    # Slide 11: Wrap Up
    slide11_repl = {
        "I built the HCMS Recruitment Suite at FIT, a full hiring platform with an AI CV-screening engine at its core, taking a candidate from CV to a digitally signed offer with a complete audit trail.":
        "I built InsightPulse at FIT: a multi-channel feedback system incorporating Voice AI and WhatsApp to centralize customer sentiment into a single analytics platform.",
        "Biggest takeaway: applying prompt engineering from my course (DA 3022) to a real product, getting an LLM to return reliable, structured, explainable output.":
        "Biggest takeaway: applying prompt engineering & structured JSON output from course DA 3022 to handle complex, multi-channel inputs and model failover configurations.",
        "Future work: embedding-based matching and a LangGraph workflow, plus a bias and fairness audit before trusting the AI unattended.":
        "Future work: implementing LangGraph for conversational multi-turn survey follow-ups, semantic clustering of user complaints, and real-time negative-sentiment email routing.",
        "[your GitHub URL]": "https://github.com/workingaditya/insightpulse",
        "[your demo URL]": "https://insightpulse.skill-pulse.io",
        "DA 377, Internship-I": "DA 378, Internship-II"
    }

    all_replacements = [
        None, # Slide 1 (done separately)
        slide2_repl,
        slide3_repl,
        slide4_repl,
        slide5_repl,
        slide6_repl,
        slide7_repl,
        slide8_repl,
        slide9_repl,
        slide10_repl,
        slide11_repl
    ]
    
    # Perform slide-by-slide replacements and design formatting
    for idx in range(1, len(prs.slides)):
        slide = prs.slides[idx]
        repls = all_replacements[idx]
        
        # --- Clean Diagram and Flowchart Leftovers ---
        shapes_to_delete = []
        
        # Slide 6 (How It Works): Delete old hiring flowchart boxes and arrows
        if idx == 5:
            # Delete Shapes 5 to 19 (which are the flowchart groups and arrows)
            for s_idx, shape in enumerate(slide.shapes):
                if 5 <= s_idx <= 19:
                    shapes_to_delete.append(shape)
                    
        # Slide 7 (AI & Voice AI): Delete old CV screening backup model tiers
        elif idx == 6:
            # Delete Shapes 5 to 9 (fallback model cards and arrows)
            for s_idx, shape in enumerate(slide.shapes):
                if 5 <= s_idx <= 9:
                    shapes_to_delete.append(shape)
                    
        # Slide 9 (CRM Integration): Delete old approval sequential flowchart
        elif idx == 8:
            # Delete Shapes 6 to 14 (flowchart boxes and arrows at bottom)
            for s_idx, shape in enumerate(slide.shapes):
                if 6 <= s_idx <= 14:
                    shapes_to_delete.append(shape)
                    
        # Delete marked diagram shapes
        for shape in shapes_to_delete:
            el = shape.element
            el.getparent().remove(el)
            
        # Apply slide-specific replacements FIRST, then global footer replacement
        if repls:
            for shape in slide.shapes:
                replace_text_recursively(shape, repls)
                
        for shape in slide.shapes:
            replace_text_recursively(shape, footer_repl)
            
        # Reposition remaining text boxes after diagram deletion to fill the screen nicely
        if idx == 5: # Slide 6
            # Reposition the main content box (formerly Shape 20)
            for shape in slide.shapes:
                if shape.has_text_frame and "Feedback flows" in shape.text_frame.text:
                    shape.left = Inches(0.75)
                    shape.top = Inches(2.00)
                    shape.width = Inches(11.90)
                    shape.height = Inches(4.20)
        elif idx == 7: # Slide 8 (academic - adjust right content column to fit well)
            for shape in slide.shapes:
                if shape.has_text_frame and "Prompt anatomy" in shape.text_frame.text:
                    shape.left = Inches(5.60)
                    shape.top = Inches(1.80)
                    shape.width = Inches(7.10)
                    shape.height = Inches(4.50)
        elif idx == 8: # Slide 9
            # Reposition the main content box (formerly Shape 5)
            for shape in slide.shapes:
                if shape.has_text_frame and "Zoho CRM Integration" in shape.text_frame.text:
                    shape.left = Inches(0.75)
                    shape.top = Inches(1.80)
                    shape.width = Inches(11.90)
                    shape.height = Inches(4.50)
            
        # Detect if there are background content Rectangle cards in this slide
        has_bg_rect = False
        for shape in slide.shapes:
            if shape.name == "Rectangle" and shape.shape_type == 1:
                top = shape.top.inches if shape.top else 0
                width = shape.width.inches if shape.width else 0
                if top > 1.5 and width < 13.0:
                    has_bg_rect = True
                    break
                
        # Design Enhancements: Card styling and paragraph formatting
        for shape in slide.shapes:
            left = shape.left.inches if shape.left else 0
            top = shape.top.inches if shape.top else 0
            width = shape.width.inches if shape.width else 0
            
            if shape.has_text_frame:
                # Format slide category headers (amber text on top banner)
                if (0.1 <= top <= 0.35) and (0.5 <= left <= 1.0) and width > 1.0:
                    for p in shape.text_frame.paragraphs:
                        p.font.name = "Trebuchet MS"
                        p.font.size = Pt(10)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(0xD9, 0x77, 0x06) # Amber
                        
                # Format main slide titles (white text on top banner)
                elif (0.35 <= top <= 0.8) and (0.5 <= left <= 1.0) and width > 8.0:
                    for p in shape.text_frame.paragraphs:
                        p.font.name = "Trebuchet MS"
                        p.font.size = Pt(24)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White
                        
                # Format slide subtitles (light-slate text on top banner, below title)
                elif (0.8 < top <= 1.25) and (0.5 <= left <= 1.0) and width > 8.0:
                    for p in shape.text_frame.paragraphs:
                        p.font.name = "Trebuchet MS"
                        p.font.size = Pt(13)
                        p.font.bold = False
                        p.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0) # Light Slate
                        
                # Format body paragraphs inside content text boxes
                else:
                    is_content_box = shape.name.startswith("•") or "Who I am" in shape.text_frame.text or "Feedback collection" in shape.text_frame.text or "This project, InsightPulse" in shape.text_frame.text or "Across the internship" in shape.text_frame.text or "Real-time feedback stream" in shape.text_frame.text or "Voice calls via VAPI" in shape.text_frame.text or "Zoho CRM OAuth" in shape.text_frame.text or "Trigger Outbound Voice AI" in shape.text_frame.text or "I built InsightPulse" in shape.text_frame.text or "Concept  →" in shape.text_frame.text or "Feedback flows" in shape.text_frame.text or "Prompt anatomy" in shape.text_frame.text
                    
                    if is_content_box:
                        format_paragraphs_in_shape(shape)
                        # If this slide doesn't have a background card shape, style this text box as a card!
                        if not has_bg_rect:
                            apply_card_styling(shape)
            
            # Apply card styling to background rectangles themselves
            elif shape.name == "Rectangle" and shape.shape_type == 1:
                if top > 1.5 and width < 13.0:
                    apply_card_styling(shape)
                    
    # Save the updated, uncorrupted presentation
    prs.save(dest_path)
    print("Presentation generated and styled successfully!")

if __name__ == "__main__":
    main()
