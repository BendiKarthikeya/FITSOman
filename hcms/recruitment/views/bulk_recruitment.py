"""
Bulk Recruitment — upload an Excel file with (Title, Vacancy) rows plus an
approval document and create Recruitment records in one shot.
"""

from django.contrib import messages
from django.shortcuts import redirect, render

from fits.decorators import login_required
from recruitment.models import Recruitment


def _parse_excel(file_obj):
    """Return (rows, error). rows = [{"title": str, "vacancy": int}, ...]."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(filename=file_obj, read_only=True, data_only=True)
    except Exception as exc:
        return [], f"Could not open Excel file: {exc}"

    ws = wb.active
    rows = []
    header = None
    for raw in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c or "").strip().lower() for c in raw]
            if "title" not in header or "vacancy" not in header:
                return [], "Excel must have a header row with 'Title' and 'Vacancy' columns."
            t_idx = header.index("title")
            v_idx = header.index("vacancy")
            continue
        if not raw or all(c in (None, "") for c in raw):
            continue
        title = str(raw[t_idx] or "").strip() if t_idx < len(raw) else ""
        try:
            vacancy = int(raw[v_idx]) if v_idx < len(raw) and raw[v_idx] not in (None, "") else 0
        except (TypeError, ValueError):
            vacancy = 0
        if not title:
            continue
        rows.append({"title": title, "vacancy": vacancy})
    return rows, None


@login_required
def bulk_recruitment(request):
    if request.method == "GET":
        return render(request, "recruitment/bulk_recruitment_form.html", {})

    excel = request.FILES.get("excel_file")
    approval = request.FILES.get("approval_document")

    if not excel:
        messages.error(request, "Please upload an Excel file.")
        return redirect("raise-recruitment")
    if not approval:
        messages.error(request, "Please attach the approval document.")
        return redirect("raise-recruitment")

    rows, err = _parse_excel(excel)
    if err:
        messages.error(request, err)
        return redirect("raise-recruitment")
    if not rows:
        messages.error(request, "No rows found in the Excel file.")
        return redirect("raise-recruitment")

    emp = getattr(request.user, "employee_get", None)
    company = getattr(emp, "company_id", None)

    from recruitment.views.views import create_employee_recruitment_approvals

    created = 0
    for row in rows:
        rec = Recruitment.objects.create(
            title=row["title"],
            vacancy=row["vacancy"] or 0,
            company_id=company,
            posting_type="internal",
            raised_by=emp,
            raised_from_employee=True,
            approval_status="pending",
            budget_document=approval,
            description=f"Bulk-imported recruitment: {row['title']}",
        )
        create_employee_recruitment_approvals(rec)
        created += 1

    messages.success(request, f"Created {created} recruitment record(s) from bulk upload.")
    return redirect("raise-recruitment")
